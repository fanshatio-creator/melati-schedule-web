from fastapi import FastAPI, APIRouter, UploadFile, File, HTTPException
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import io
import re
import uuid
import logging
import calendar
from pathlib import Path
from fastapi.responses import Response
from pydantic import BaseModel
from typing import List, Optional
from datetime import datetime, timezone, timedelta, date
import pandas as pd

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

RUANGAN = [
    {"id": "ruang-tu", "name": "Ruangan TU (Klaster Manajemen)", "code": "TU-KM", "category": "Manajemen & TU"},
    {"id": "ruang-klaster-3", "name": "Ruang Klaster 3", "code": "KL-03", "category": "Pelayanan Usia Dewasa & Lansia"},
    {"id": "ruang-klaster-2-ibu", "name": "Ruang Klaster 2 Ibu", "code": "KL2-IBU", "category": "Pelayanan Kesehatan Ibu"},
    {"id": "ruang-klaster-2-anak", "name": "Ruang Klaster 2 Anak", "code": "KL2-ANK", "category": "Pelayanan Kesehatan Anak"},
    {"id": "ruang-imunisasi", "name": "Ruang Imunisasi", "code": "RM-IMN", "category": "Pencegahan & Imunisasi"},
    {"id": "ruang-klaster-gigi", "name": "Ruang Klaster Gigi dan Mulut", "code": "KL-GGI", "category": "Kesehatan Gigi & Mulut"},
    {"id": "meja-skrining", "name": "Meja Skrining", "code": "MJ-SKR", "category": "Triase & Skrining Awal"},
    {"id": "ruang-tindakan", "name": "Ruang Tindakan", "code": "RM-TND", "category": "Gawat Darurat & Tindakan"},
    {"id": "ruang-laboratorium", "name": "Ruang Laboratorium", "code": "RM-LAB", "category": "Penunjang Medis"},
    {"id": "ruang-farmasi", "name": "Ruang Farmasi", "code": "RM-FRM", "category": "Kefarmasian & Obat"},
    {"id": "loket", "name": "Loket", "code": "LKT-01", "category": "Pelayanan Informasi & Administrasi"},
    {"id": "pendaftaran-antrian", "name": "Pendaftaran Antrian", "code": "PDF-ANT", "category": "Pendaftaran Pasien"},
    {"id": "klaster-4", "name": "Klaster 4", "code": "KL-04", "category": "Penanggulangan Penyakit Menular"},
]

RUANGAN_MAP = {r["id"]: r["name"] for r in RUANGAN}

BULAN_ID = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
            "Agustus", "September", "Oktober", "November", "Desember"]

STATUS_LABEL_ID = {"menunggu": "Menunggu Persetujuan", "disetujui": "Disetujui", "ditolak": "Ditolak"}

PEGAWAI_ALIASES = {
    "nip": ["nip", "nik", "no induk"],
    "nama": ["nama"],
    "jabatan": ["jabatan", "profesi"],
    "ruangan": ["ruangan", "ruang", "unit kerja", "unit"],
    "telepon": ["telepon", "telp", "hp", "wa", "kontak"],
    "status": ["status"],
}

JADWAL_ALIASES = {
    "nama_kegiatan": ["kegiatan", "agenda", "acara", "program"],
    "tanggal_mulai": ["tanggal mulai", "tgl mulai", "mulai", "tanggal", "tgl"],
    "tanggal_selesai": ["tanggal selesai", "tgl selesai", "selesai", "berakhir", "sampai"],
    "lokasi": ["lokasi", "tempat", "faskes"],
    "pegawai": ["pegawai", "petugas", "personil", "penanggung jawab", "pelaksana"],
    "koordinator": ["koordinator", "pengisi", "pj program", "koordinator program", "pelapor"],
    "keterangan": ["keterangan", "catatan", "output"],
}

DATE_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y/%m/%d", "%d %B %Y", "%d %b %Y"]


def norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def parse_date_value(v):
    if v is None:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    s = str(v).strip()
    if not s or s.lower() == "nan":
        return None
    for f in DATE_FORMATS:
        try:
            return datetime.strptime(s, f).date().isoformat()
        except ValueError:
            pass
    try:
        return pd.to_datetime(s, dayfirst=True).date().isoformat()
    except Exception:
        return None


def match_ruangan_id(name):
    n = norm(name)
    if not n:
        return ""
    for r in RUANGAN:
        if n == norm(r["id"]) or n == norm(r["code"]):
            return r["id"]
    for r in RUANGAN:
        rn = norm(r["name"])
        if n in rn or rn in n:
            return r["id"]
    aliases = {"tu": "ruang-tu", "imunisasi": "ruang-imunisasi", "gigi": "ruang-klaster-gigi",
               "lab": "ruang-laboratorium", "laboratorium": "ruang-laboratorium",
               "farmasi": "ruang-farmasi", "skrining": "meja-skrining", "tindakan": "ruang-tindakan",
               "pendaftaran": "pendaftaran-antrian", "antrian": "pendaftaran-antrian",
               "klaster3": "ruang-klaster-3", "klaster4": "klaster-4",
               "klaster2ibu": "ruang-klaster-2-ibu", "klaster2anak": "ruang-klaster-2-anak"}
    return aliases.get(n, "")


def read_table(filename, content):
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext in ("xlsx", "xls"):
        df = pd.read_excel(io.BytesIO(content), header=None)
        return df.fillna("").astype(str).values.tolist()
    if ext == "csv":
        df = pd.read_csv(io.BytesIO(content), header=None)
        return df.fillna("").astype(str).values.tolist()
    if ext == "docx":
        from docx import Document
        doc = Document(io.BytesIO(content))
        rows = []
        for t in doc.tables:
            for r in t.rows:
                rows.append([c.text.strip() for c in r.cells])
        return rows
    if ext == "pdf":
        import pdfplumber
        rows = []
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                tables = page.extract_tables()
                if tables:
                    for table in tables:
                        rows.extend([[(c or "").strip() for c in row] for row in table])
                else:
                    words = page.extract_words()
                    lines = {}
                    for w in words:
                        key = None
                        for k in lines:
                            if abs(k - w["top"]) < 3:
                                key = k
                                break
                        if key is None:
                            key = w["top"]
                        lines.setdefault(key, []).append(w)
                    for top in sorted(lines):
                        ws = sorted(lines[top], key=lambda x: x["x0"])
                        cells, cur = [], []
                        prev_x1 = None
                        for w in ws:
                            if prev_x1 is not None and (w["x0"] - prev_x1) > 8:
                                cells.append(" ".join(cur))
                                cur = []
                            cur.append(w["text"])
                            prev_x1 = w["x1"]
                        if cur:
                            cells.append(" ".join(cur))
                        if cells:
                            rows.append(cells)
        return rows
    raise HTTPException(400, "Format file tidak didukung. Gunakan .xlsx, .xls, .csv, .pdf, atau .docx")


def rows_to_dicts(table, aliases):
    header_idx, colmap = None, {}
    for i, row in enumerate(table[:15]):
        cmap = {}
        for idx, cell in enumerate(row):
            n = norm(cell)
            if not n:
                continue
            for key, keys in aliases.items():
                if key in cmap:
                    continue
                if any(norm(k) and (norm(k) in n) for k in keys):
                    cmap[key] = idx
                    break
        if len(cmap) >= 2:
            header_idx, colmap = i, cmap
            break
    if header_idx is None:
        raise HTTPException(400, "Tidak menemukan baris header tabel di file. Pastikan file berisi tabel dengan judul kolom.")
    out = []
    for row in table[header_idx + 1:]:
        if not any(str(c).strip() for c in row):
            continue
        item = {}
        for key, idx in colmap.items():
            item[key] = str(row[idx]).strip() if idx < len(row) else ""
        out.append(item)
    return out


async def find_conflicts(pegawai_ids, t_mulai, t_selesai, exclude_id=None):
    if not pegawai_ids:
        return []
    query = {
        "status": {"$ne": "ditolak"},
        "tanggal_mulai": {"$lte": t_selesai},
        "tanggal_selesai": {"$gte": t_mulai},
        "pegawai_ids": {"$in": list(pegawai_ids)},
    }
    if exclude_id:
        query["id"] = {"$ne": exclude_id}
    docs = await db.jadwal.find(query, {"_id": 0}).to_list(1000)
    pids = set(pegawai_ids)
    all_p = await db.pegawai.find({"id": {"$in": list(pids)}}, {"_id": 0}).to_list(1000)
    pmap = {p["id"]: p for p in all_p}
    conflicts = []
    for d in docs:
        overlap = pids & set(d.get("pegawai_ids", []))
        if overlap:
            conflicts.append({
                "jadwal_id": d["id"],
                "nama_kegiatan": d["nama_kegiatan"],
                "tanggal_mulai": d["tanggal_mulai"],
                "tanggal_selesai": d["tanggal_selesai"],
                "lokasi": d.get("lokasi", ""),
                "status": d.get("status", ""),
                "pegawai": [pmap[pid]["nama"] for pid in overlap if pid in pmap],
            })
    return conflicts


class PegawaiInput(BaseModel):
    nip: str = ""
    nama: str
    jabatan: str = ""
    ruangan_id: str = ""
    telepon: str = ""
    status_kepegawaian: str = ""


class BulkPegawaiRow(BaseModel):
    nip: str = ""
    nama: str = ""
    jabatan: str = ""
    ruangan: str = ""
    telepon: str = ""
    status: str = ""


class BulkPegawaiInput(BaseModel):
    rows: List[BulkPegawaiRow]


class JadwalInput(BaseModel):
    nama_kegiatan: str
    tanggal_mulai: str
    tanggal_selesai: str
    lokasi: str = ""
    koordinator: str = ""
    keterangan: str = ""
    pegawai_ids: List[str] = []


class BulkJadwalRow(BaseModel):
    nama_kegiatan: str = ""
    tanggal_mulai: str = ""
    tanggal_selesai: str = ""
    lokasi: str = ""
    pegawai: str = ""
    pegawai_ids: List[str] = []
    koordinator: str = ""
    keterangan: str = ""


class BulkJadwalInput(BaseModel):
    rows: List[BulkJadwalRow]


class ApprovalInput(BaseModel):
    aksi: str
    approver: str
    catatan: Optional[str] = ""


PIN_ROLES = ["Kepala TU", "Kepala Puskesmas"]
DEFAULT_PINS = {"Kepala TU": "1234", "Kepala Puskesmas": "4321"}


class VerifyPinInput(BaseModel):
    role: str
    pin: str


class ChangePinInput(BaseModel):
    role: str
    pin_lama: str
    pin_baru: str


async def get_pins():
    doc = await db.app_settings.find_one({"key": "pins"}, {"_id": 0})
    if not doc:
        doc = {"key": "pins", "pins": dict(DEFAULT_PINS)}
        await db.app_settings.insert_one(dict(doc))
    return doc.get("pins", dict(DEFAULT_PINS))


@api_router.get("/")
async def root():
    return {"message": "API Jadwal Puskesmas aktif"}


@api_router.get("/ruangan")
async def get_ruangan():
    return RUANGAN


# ---------- PEGAWAI ----------
@api_router.get("/pegawai")
async def list_pegawai():
    return await db.pegawai.find({}, {"_id": 0}).sort("nama", 1).to_list(2000)


@api_router.post("/pegawai")
async def create_pegawai(input: PegawaiInput):
    if input.nip:
        existing = await db.pegawai.find_one({"nip": input.nip}, {"_id": 0})
        if existing:
            raise HTTPException(409, f"NIP {input.nip} sudah terdaftar atas nama {existing['nama']}")
    doc = input.model_dump()
    doc["id"] = str(uuid.uuid4())
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    await db.pegawai.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.put("/pegawai/{pid}")
async def update_pegawai(pid: str, input: PegawaiInput):
    if input.nip:
        existing = await db.pegawai.find_one({"nip": input.nip, "id": {"$ne": pid}}, {"_id": 0})
        if existing:
            raise HTTPException(409, f"NIP {input.nip} sudah terdaftar atas nama {existing['nama']}")
    res = await db.pegawai.update_one({"id": pid}, {"$set": input.model_dump()})
    if res.matched_count == 0:
        raise HTTPException(404, "Pegawai tidak ditemukan")
    return await db.pegawai.find_one({"id": pid}, {"_id": 0})


@api_router.delete("/pegawai/{pid}")
async def delete_pegawai(pid: str):
    res = await db.pegawai.delete_one({"id": pid})
    if res.deleted_count == 0:
        raise HTTPException(404, "Pegawai tidak ditemukan")
    await db.jadwal.update_many({}, {"$pull": {"pegawai_ids": pid}})
    return {"ok": True}


@api_router.post("/pegawai/import/parse")
async def parse_pegawai_file(file: UploadFile = File(...)):
    content = await file.read()
    table = read_table(file.filename, content)
    rows = rows_to_dicts(table, PEGAWAI_ALIASES)
    return {"rows": rows, "count": len(rows)}


@api_router.post("/pegawai/bulk")
async def bulk_pegawai(input: BulkPegawaiInput):
    added, skipped = [], []
    seen_nips = set()
    for row in input.rows:
        nama = row.nama.strip()
        if not nama:
            skipped.append({"nama": "-", "reason": "Nama kosong"})
            continue
        nip = row.nip.strip()
        if nip and nip in seen_nips:
            skipped.append({"nama": nama, "reason": f"NIP {nip} duplikat di dalam file"})
            continue
        if nip:
            existing = await db.pegawai.find_one({"nip": nip}, {"_id": 0})
            if existing:
                skipped.append({"nama": nama, "reason": f"NIP {nip} sudah terdaftar ({existing['nama']})"})
                continue
        doc = {
            "id": str(uuid.uuid4()),
            "nip": nip,
            "nama": nama,
            "jabatan": row.jabatan.strip(),
            "ruangan_id": match_ruangan_id(row.ruangan),
            "telepon": row.telepon.strip(),
            "status_kepegawaian": row.status.strip(),
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        await db.pegawai.insert_one(doc)
        if nip:
            seen_nips.add(nip)
        added.append(nama)
    return {"added": len(added), "added_names": added, "skipped": skipped}


# ---------- JADWAL ----------
@api_router.get("/jadwal")
async def list_jadwal():
    docs = await db.jadwal.find({}, {"_id": 0}).sort("tanggal_mulai", -1).to_list(2000)
    pids = list({pid for d in docs for pid in d.get("pegawai_ids", [])})
    plist = await db.pegawai.find({"id": {"$in": pids}}, {"_id": 0}).to_list(2000) if pids else []
    pmap = {p["id"]: p for p in plist}
    for d in docs:
        d["pegawai"] = [{"id": pid, "nama": pmap[pid]["nama"], "nip": pmap[pid].get("nip", "")}
                        for pid in d.get("pegawai_ids", []) if pid in pmap]
    return docs


@api_router.get("/jadwal/conflicts/check")
async def check_conflicts(pegawai_ids: str = "", tanggal_mulai: str = "", tanggal_selesai: str = "", exclude_id: str = ""):
    pids = [p for p in pegawai_ids.split(",") if p]
    if not pids or not tanggal_mulai or not tanggal_selesai:
        return {"conflicts": []}
    conflicts = await find_conflicts(pids, tanggal_mulai, tanggal_selesai, exclude_id or None)
    return {"conflicts": conflicts}


@api_router.post("/jadwal")
async def create_jadwal(input: JadwalInput):
    if input.tanggal_selesai < input.tanggal_mulai:
        raise HTTPException(400, "Tanggal selesai tidak boleh lebih awal dari tanggal mulai")
    conflicts = await find_conflicts(input.pegawai_ids, input.tanggal_mulai, input.tanggal_selesai)
    if conflicts:
        raise HTTPException(409, {"message": "Jadwal bentrok terdeteksi. Penyimpanan diblokir.", "conflicts": conflicts})
    doc = input.model_dump()
    doc["id"] = str(uuid.uuid4())
    doc["status"] = "menunggu"
    doc["disetujui_oleh"] = ""
    doc["catatan_approval"] = ""
    doc["waktu_approval"] = ""
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    await db.jadwal.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.put("/jadwal/{jid}")
async def update_jadwal(jid: str, input: JadwalInput):
    existing = await db.jadwal.find_one({"id": jid}, {"_id": 0})
    if not existing:
        raise HTTPException(404, "Jadwal tidak ditemukan")
    if input.tanggal_selesai < input.tanggal_mulai:
        raise HTTPException(400, "Tanggal selesai tidak boleh lebih awal dari tanggal mulai")
    conflicts = await find_conflicts(input.pegawai_ids, input.tanggal_mulai, input.tanggal_selesai, exclude_id=jid)
    if conflicts:
        raise HTTPException(409, {"message": "Jadwal bentrok terdeteksi. Penyimpanan diblokir.", "conflicts": conflicts})
    update = input.model_dump()
    update["status"] = "menunggu"
    update["disetujui_oleh"] = ""
    update["catatan_approval"] = ""
    update["waktu_approval"] = ""
    await db.jadwal.update_one({"id": jid}, {"$set": update})
    return await db.jadwal.find_one({"id": jid}, {"_id": 0})


@api_router.delete("/jadwal/{jid}")
async def delete_jadwal(jid: str):
    res = await db.jadwal.delete_one({"id": jid})
    if res.deleted_count == 0:
        raise HTTPException(404, "Jadwal tidak ditemukan")
    return {"ok": True}


@api_router.post("/jadwal/{jid}/approval")
async def approval_jadwal(jid: str, input: ApprovalInput):
    if input.aksi not in ("setujui", "tolak"):
        raise HTTPException(400, "Aksi harus 'setujui' atau 'tolak'")
    if input.approver not in ("Kepala TU", "Kepala Puskesmas"):
        raise HTTPException(403, "Hanya Kepala TU atau Kepala Puskesmas yang dapat memberi persetujuan")
    existing = await db.jadwal.find_one({"id": jid}, {"_id": 0})
    if not existing:
        raise HTTPException(404, "Jadwal tidak ditemukan")
    status = "disetujui" if input.aksi == "setujui" else "ditolak"
    await db.jadwal.update_one({"id": jid}, {"$set": {
        "status": status,
        "disetujui_oleh": input.approver,
        "catatan_approval": input.catatan or "",
        "waktu_approval": datetime.now(timezone.utc).isoformat(),
    }})
    return await db.jadwal.find_one({"id": jid}, {"_id": 0})


# ---------- AUTH / PIN ----------
@api_router.post("/auth/verify-pin")
async def verify_pin(input: VerifyPinInput):
    if input.role not in PIN_ROLES:
        return {"ok": True}
    pins = await get_pins()
    if input.pin == pins.get(input.role, DEFAULT_PINS.get(input.role)):
        return {"ok": True}
    raise HTTPException(401, "PIN salah")


@api_router.post("/auth/change-pin")
async def change_pin(input: ChangePinInput):
    if input.role not in PIN_ROLES:
        raise HTTPException(400, "Role ini tidak menggunakan PIN")
    if len(input.pin_baru) < 4 or not input.pin_baru.isdigit():
        raise HTTPException(400, "PIN baru harus berupa minimal 4 digit angka")
    pins = await get_pins()
    if input.pin_lama != pins.get(input.role, DEFAULT_PINS.get(input.role)):
        raise HTTPException(401, "PIN lama salah")
    pins[input.role] = input.pin_baru
    await db.app_settings.update_one({"key": "pins"}, {"$set": {"pins": pins}}, upsert=True)
    return {"ok": True, "message": "PIN berhasil diperbarui"}


# ---------- MATRIKS / SINKRONISASI ----------
@api_router.get("/jadwal/matrix")
async def jadwal_matrix(year: int, month: int):
    if month < 1 or month > 12:
        raise HTTPException(400, "Bulan tidak valid")
    ndays = calendar.monthrange(year, month)[1]
    month_start = date(year, month, 1).isoformat()
    month_end = date(year, month, ndays).isoformat()
    hari = ["Sen", "Sel", "Rab", "Kam", "Jum", "Sab", "Min"]
    dates = []
    for d in range(1, ndays + 1):
        dt = date(year, month, d)
        dates.append({
            "iso": dt.isoformat(),
            "day": d,
            "weekday": hari[dt.weekday()],
            "is_weekend": dt.weekday() >= 5,
        })

    docs = await db.jadwal.find({
        "status": {"$ne": "ditolak"},
        "tanggal_mulai": {"$lte": month_end},
        "tanggal_selesai": {"$gte": month_start},
    }, {"_id": 0}).sort("tanggal_mulai", 1).to_list(2000)

    pids = list({pid for d in docs for pid in d.get("pegawai_ids", [])})
    plist = await db.pegawai.find({"id": {"$in": pids}}, {"_id": 0}).to_list(2000) if pids else []
    pmap = {p["id"]: p for p in plist}

    def in_range(iso_day, jd):
        return jd["tanggal_mulai"] <= iso_day <= jd["tanggal_selesai"]

    rows = []
    conflicts = []
    for pid in pids:
        p = pmap.get(pid)
        if not p:
            continue
        cells = {}
        for dt in dates:
            iso_day = dt["iso"]
            acts = [{
                "jadwal_id": j["id"],
                "nama_kegiatan": j["nama_kegiatan"],
                "lokasi": j.get("lokasi", ""),
                "status": j.get("status", ""),
            } for j in docs if pid in j.get("pegawai_ids", []) and in_range(iso_day, j)]
            if acts:
                cells[iso_day] = acts
                if len(acts) > 1:
                    conflicts.append({
                        "tanggal": iso_day,
                        "pegawai_id": pid,
                        "pegawai_nama": p["nama"],
                        "kegiatan": [a["nama_kegiatan"] for a in acts],
                    })
        if cells:
            rows.append({
                "id": pid,
                "nama": p["nama"],
                "jabatan": p.get("jabatan", ""),
                "cells": cells,
                "total": sum(len(v) for v in cells.values()),
            })
    rows.sort(key=lambda r: r["nama"])

    kegiatan = []
    for j in docs:
        kegiatan.append({
            "id": j["id"],
            "nama_kegiatan": j["nama_kegiatan"],
            "lokasi": j.get("lokasi", ""),
            "status": j.get("status", ""),
            "koordinator": j.get("koordinator", ""),
            "tanggal_mulai": j["tanggal_mulai"],
            "tanggal_selesai": j["tanggal_selesai"],
            "pegawai": [{"id": pid, "nama": pmap[pid]["nama"]} for pid in j.get("pegawai_ids", []) if pid in pmap],
        })

    return {
        "year": year,
        "month": month,
        "dates": dates,
        "rows": rows,
        "kegiatan": kegiatan,
        "conflicts": conflicts,
        "total_pegawai_terjadwal": len(rows),
        "total_kegiatan": len(kegiatan),
    }


@api_router.post("/jadwal/import/parse")
async def parse_jadwal_file(file: UploadFile = File(...)):
    content = await file.read()
    table = read_table(file.filename, content)
    rows = rows_to_dicts(table, JADWAL_ALIASES)
    all_pegawai = await db.pegawai.find({}, {"_id": 0}).to_list(2000)
    by_name = {norm(p["nama"]): p for p in all_pegawai}
    by_nip = {norm(p.get("nip", "")): p for p in all_pegawai if p.get("nip")}
    out = []
    for r in rows:
        mulai = parse_date_value(r.get("tanggal_mulai"))
        selesai = parse_date_value(r.get("tanggal_selesai")) or mulai
        pegawai_raw = r.get("pegawai", "")
        pids, unmatched = [], []
        for nm in re.split(r"[;\n]", pegawai_raw or ""):
            nm = nm.strip()
            if not nm:
                continue
            p = by_name.get(norm(nm)) or by_nip.get(norm(nm))
            if p and p["id"] not in pids:
                pids.append(p["id"])
            elif not p:
                unmatched.append(nm)
        out.append({
            "nama_kegiatan": r.get("nama_kegiatan", ""),
            "tanggal_mulai": mulai or r.get("tanggal_mulai", ""),
            "tanggal_selesai": selesai or r.get("tanggal_selesai", "") or (mulai or ""),
            "lokasi": r.get("lokasi", ""),
            "pegawai": pegawai_raw,
            "pegawai_ids": pids,
            "pegawai_unmatched": unmatched,
            "koordinator": r.get("koordinator", ""),
            "keterangan": r.get("keterangan", ""),
        })
    return {"rows": out, "count": len(out)}


@api_router.post("/jadwal/bulk")
async def bulk_jadwal(input: BulkJadwalInput):
    all_pegawai = await db.pegawai.find({}, {"_id": 0}).to_list(2000)
    by_name = {norm(p["nama"]): p for p in all_pegawai}
    by_nip = {norm(p.get("nip", "")): p for p in all_pegawai if p.get("nip")}
    id_set = {p["id"] for p in all_pegawai}
    added, skipped = [], []
    accepted_ranges = {}
    for row in input.rows:
        label = row.nama_kegiatan.strip() or "-"
        if not row.nama_kegiatan.strip():
            skipped.append({"nama_kegiatan": "-", "reason": "Nama kegiatan kosong"})
            continue
        mulai = parse_date_value(row.tanggal_mulai)
        selesai = parse_date_value(row.tanggal_selesai) or mulai
        if not mulai:
            skipped.append({"nama_kegiatan": label, "reason": f"Tanggal tidak valid: '{row.tanggal_mulai}'"})
            continue
        if selesai < mulai:
            skipped.append({"nama_kegiatan": label, "reason": "Tanggal selesai lebih awal dari tanggal mulai"})
            continue
        pids, unknown = [], []
        if row.pegawai_ids:
            pids = [pid for pid in row.pegawai_ids if pid in id_set]
        else:
            for nm in re.split(r"[;\n]", row.pegawai or ""):
                nm = nm.strip()
                if not nm:
                    continue
                p = by_name.get(norm(nm)) or by_nip.get(norm(nm))
                if p:
                    pids.append(p["id"])
                else:
                    unknown.append(nm)
            if row.pegawai.strip() and not pids:
                skipped.append({"nama_kegiatan": label, "reason": f"Pegawai tidak ditemukan: {', '.join(unknown)}"})
                continue
        conflicts = await find_conflicts(pids, mulai, selesai)
        if conflicts:
            names = sorted({n for c in conflicts for n in c["pegawai"]})
            skipped.append({"nama_kegiatan": label, "reason": f"Bentrok dengan jadwal lain: {', '.join(names)}"})
            continue
        file_clash = False
        for pid in pids:
            for (a, b) in accepted_ranges.get(pid, []):
                if mulai <= b and a <= selesai:
                    file_clash = True
        if file_clash:
            skipped.append({"nama_kegiatan": label, "reason": "Bentrok dengan baris lain di file yang sama"})
            continue
        doc = {
            "id": str(uuid.uuid4()),
            "nama_kegiatan": row.nama_kegiatan.strip(),
            "tanggal_mulai": mulai,
            "tanggal_selesai": selesai,
            "lokasi": row.lokasi.strip(),
            "koordinator": row.koordinator.strip(),
            "keterangan": row.keterangan.strip(),
            "pegawai_ids": pids,
            "status": "menunggu",
            "disetujui_oleh": "",
            "catatan_approval": "",
            "waktu_approval": "",
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        await db.jadwal.insert_one(doc)
        for pid in pids:
            accepted_ranges.setdefault(pid, []).append((mulai, selesai))
        added.append(label)
    return {"added": len(added), "added_names": added, "skipped": skipped}


# ---------- EXPORT ----------
async def _jadwal_for_month(year: int, month: int):
    first = date(year, month, 1)
    last = date(year, month, calendar.monthrange(year, month)[1])
    docs = await db.jadwal.find({
        "tanggal_mulai": {"$lte": last.isoformat()},
        "tanggal_selesai": {"$gte": first.isoformat()},
    }, {"_id": 0}).sort("tanggal_mulai", 1).to_list(2000)
    pids = list({pid for d in docs for pid in d.get("pegawai_ids", [])})
    plist = await db.pegawai.find({"id": {"$in": pids}}, {"_id": 0}).to_list(2000) if pids else []
    pmap = {p["id"]: p for p in plist}
    for d in docs:
        d["pegawai_nama"] = ", ".join(pmap[pid]["nama"] for pid in d.get("pegawai_ids", []) if pid in pmap)
    return docs, first, last


def _fmt_range(mulai, selesai):
    if mulai == selesai:
        return mulai
    return f"{mulai} s/d {selesai}"


@api_router.get("/jadwal/export/excel")
async def export_jadwal_excel(year: int, month: int):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    docs, first, last = await _jadwal_for_month(year, month)
    judul = f"Jadwal Kegiatan Luar Puskesmas — {BULAN_ID[month - 1]} {year}"

    wb = Workbook()
    ws = wb.active
    ws.title = f"{BULAN_ID[month - 1]} {year}"[:31]
    headers = ["No", "Nama Kegiatan", "Tanggal", "Lokasi / Faskes", "Pegawai Ditugaskan", "Koordinator Program", "Status", "Keterangan"]

    ws.merge_cells("A1:H1")
    ws["A1"] = judul
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="047857")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    head_fill = PatternFill("solid", fgColor="D1FAE5")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.font = Font(bold=True, color="065F46")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for i, d in enumerate(docs, start=1):
        row = [
            i,
            d.get("nama_kegiatan", ""),
            _fmt_range(d.get("tanggal_mulai", ""), d.get("tanggal_selesai", "")),
            d.get("lokasi", ""),
            d.get("pegawai_nama", ""),
            d.get("koordinator", ""),
            STATUS_LABEL_ID.get(d.get("status", ""), d.get("status", "")),
            d.get("keterangan", ""),
        ]
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=2 + i, column=c, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if c in (1, 3, 7) else "left")
            cell.border = border

    if not docs:
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=8)
        ws.cell(row=3, column=1, value="Tidak ada jadwal pada bulan ini").alignment = Alignment(horizontal="center")

    widths = [5, 32, 22, 24, 30, 24, 20, 28]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + c)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"jadwal-{year}-{month:02d}.xlsx"
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@api_router.get("/jadwal/export/pdf")
async def export_jadwal_pdf(year: int, month: int):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    docs, first, last = await _jadwal_for_month(year, month)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm,
                            rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], fontSize=15, textColor=colors.HexColor("#065F46"))
    sub_style = ParagraphStyle("s", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#64748B"))
    cell_style = ParagraphStyle("c", parent=styles["Normal"], fontSize=8, leading=10)
    head_style = ParagraphStyle("h", parent=styles["Normal"], fontSize=8, leading=10,
                                textColor=colors.HexColor("#065F46"), fontName="Helvetica-Bold")

    elems = [
        Paragraph("Jadwal Kegiatan Luar Puskesmas", title_style),
        Paragraph(f"Periode: {BULAN_ID[month - 1]} {year}", sub_style),
        Spacer(1, 8),
    ]

    headers = ["No", "Nama Kegiatan", "Tanggal", "Lokasi", "Pegawai", "Koordinator", "Status", "Keterangan"]
    data = [[Paragraph(h, head_style) for h in headers]]
    for i, d in enumerate(docs, start=1):
        data.append([
            Paragraph(str(i), cell_style),
            Paragraph(d.get("nama_kegiatan", ""), cell_style),
            Paragraph(_fmt_range(d.get("tanggal_mulai", ""), d.get("tanggal_selesai", "")), cell_style),
            Paragraph(d.get("lokasi", "") or "-", cell_style),
            Paragraph(d.get("pegawai_nama", "") or "-", cell_style),
            Paragraph(d.get("koordinator", "") or "-", cell_style),
            Paragraph(STATUS_LABEL_ID.get(d.get("status", ""), d.get("status", "")), cell_style),
            Paragraph(d.get("keterangan", "") or "-", cell_style),
        ])
    if not docs:
        data.append([Paragraph("Tidak ada jadwal pada bulan ini", cell_style)] + ["" for _ in range(7)])

    col_widths = [10 * mm, 46 * mm, 30 * mm, 34 * mm, 46 * mm, 30 * mm, 26 * mm, 37 * mm]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D1FAE5")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    if not docs:
        table.setStyle(TableStyle([("SPAN", (0, 1), (-1, 1)), ("ALIGN", (0, 1), (-1, 1), "CENTER")]))
    elems.append(table)
    doc.build(elems)
    buf.seek(0)
    fname = f"jadwal-{year}-{month:02d}.pdf"
    return Response(
        content=buf.read(),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@api_router.get("/jadwal/template")
async def download_jadwal_template():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    pegawai = await db.pegawai.find({}, {"_id": 0}).sort("nama", 1).to_list(2000)
    nama_list = [p["nama"] for p in pegawai]

    wb = Workbook()
    ws = wb.active
    ws.title = "Jadwal"
    headers = ["Nama Kegiatan", "Tanggal Mulai", "Tanggal Selesai", "Lokasi", "Pegawai", "Koordinator", "Keterangan"]

    head_fill = PatternFill("solid", fgColor="047857")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 24

    ex1 = nama_list[0] if nama_list else "Nama Pegawai"
    ex_multi = "; ".join(nama_list[1:4]) if len(nama_list) >= 4 else "; ".join(nama_list[:2]) if len(nama_list) >= 2 else ex1
    examples = [
        ["Posyandu Balita Kelurahan Sukamaju", "19/09/2026", "19/09/2026", "Balai RW 05", ex_multi, ex1, "Penimbangan & imunisasi dasar"],
        ["Puskesmas Keliling Desa Binaan", "22/09/2026", "23/09/2026", "Desa Binaan", "; ".join(nama_list[:2]) if len(nama_list) >= 2 else ex1, ex1, "Pemeriksaan umum gratis"],
    ]
    for i, ex in enumerate(examples, start=2):
        for c, val in enumerate(ex, start=1):
            cell = ws.cell(row=i, column=c, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border

    widths = [34, 15, 15, 24, 46, 26, 34]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + c)].width = w

    note = ("PETUNJUK PENGISIAN:\n"
            "• Satu baris = satu kegiatan.\n"
            "• Kolom 'Pegawai' BOLEH LEBIH DARI SATU petugas — pisahkan setiap nama dengan TITIK KOMA (;). "
            "Contoh: \"Nama A; Nama B; Nama C\". (Jangan gunakan koma karena banyak nama mengandung koma, mis. 'Bd. Dewi, S.Tr.Keb').\n"
            "• Setiap nama harus PERSIS sama dengan daftar di sheet 'Daftar Pegawai' (salin-tempel untuk menghindari salah ketik).\n"
            "• Tanggal format DD/MM/YYYY (mis. 19/09/2026).\n"
            "• 'Koordinator' = nama pengisi/penanggung jawab program.\n"
            "• Hapus 2 baris contoh di atas sebelum mengunggah.")
    ncell = ws.cell(row=len(examples) + 3, column=1, value=note)
    ncell.font = Font(italic=True, color="B91C1C", size=9)
    ncell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=len(examples) + 3, start_column=1, end_row=len(examples) + 3, end_column=7)
    ws.row_dimensions[len(examples) + 3].height = 108

    # Reference sheet with valid employee names
    ref = wb.create_sheet("Daftar Pegawai")
    ref_headers = ["Nama Pegawai", "NIP", "Jabatan"]
    for c, h in enumerate(ref_headers, start=1):
        cell = ref.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    for i, p in enumerate(pegawai, start=2):
        ref.cell(row=i, column=1, value=p["nama"]).border = border
        ref.cell(row=i, column=2, value=p.get("nip", "")).border = border
        ref.cell(row=i, column=3, value=p.get("jabatan", "")).border = border
    ref.column_dimensions["A"].width = 40
    ref.column_dimensions["B"].width = 26
    ref.column_dimensions["C"].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="template-import-jadwal.xlsx"'},
    )


# ---------- DASHBOARD ----------
@api_router.get("/dashboard")
async def dashboard(tanggal: str = ""):
    tgl = tanggal or datetime.now(timezone.utc).date().isoformat()
    pegawai = await db.pegawai.find({}, {"_id": 0}).sort("nama", 1).to_list(2000)
    jadwal_hari_ini = await db.jadwal.find({
        "status": {"$ne": "ditolak"},
        "tanggal_mulai": {"$lte": tgl},
        "tanggal_selesai": {"$gte": tgl},
    }, {"_id": 0}).to_list(2000)
    luar_map = {}
    for j in jadwal_hari_ini:
        for pid in j.get("pegawai_ids", []):
            luar_map.setdefault(pid, []).append({
                "kegiatan": j["nama_kegiatan"],
                "lokasi": j.get("lokasi", ""),
                "status": j["status"],
                "tanggal_mulai": j["tanggal_mulai"],
                "tanggal_selesai": j["tanggal_selesai"],
            })
    rooms = []
    di_dalam_total, di_luar_total = 0, 0
    tanpa_ruangan = {"di_dalam": [], "di_luar": []}
    per_room = {r["id"]: {"di_dalam": [], "di_luar": []} for r in RUANGAN}
    for p in pegawai:
        info = {"id": p["id"], "nip": p.get("nip", ""), "nama": p["nama"], "jabatan": p.get("jabatan", "")}
        rid = p.get("ruangan_id", "")
        bucket = per_room.get(rid)
        if pid_luar := luar_map.get(p["id"]):
            entry = {**info, "kegiatan_luar": pid_luar}
            (bucket["di_luar"] if bucket else tanpa_ruangan["di_luar"]).append(entry)
            di_luar_total += 1
        else:
            (bucket["di_dalam"] if bucket else tanpa_ruangan["di_dalam"]).append(info)
            di_dalam_total += 1
    for r in RUANGAN:
        rooms.append({**r, **per_room[r["id"]]})
    menunggu = await db.jadwal.count_documents({"status": "menunggu"})
    return {
        "tanggal": tgl,
        "total_pegawai": len(pegawai),
        "di_dalam": di_dalam_total,
        "di_luar": di_luar_total,
        "menunggu_approval": menunggu,
        "ruangan": rooms,
        "tanpa_ruangan": tanpa_ruangan,
    }


# ---------- SEED ----------
SAMPLE_PEGAWAI = [
    ("19850412 201001 1 001", "dr. Ratna Wulandari", "Dokter Umum", "ruang-klaster-3", "0812-3456-7801", "PNS"),
    ("19900315 201501 2 002", "Ns. Budi Santoso, S.Kep", "Perawat", "ruang-klaster-3", "0812-3456-7802", "PNS"),
    ("19871120 201201 1 003", "Bd. Siti Aminah, S.Tr.Keb", "Bidan", "ruang-klaster-2-ibu", "0812-3456-7803", "PNS"),
    ("19920508 201701 2 004", "Bd. Dewi Lestari, S.Tr.Keb", "Bidan", "ruang-klaster-2-anak", "0812-3456-7804", "PPPK"),
    ("19880214 201301 1 005", "Ns. Agus Prasetyo, S.Kep", "Perawat", "ruang-imunisasi", "0812-3456-7805", "PNS"),
    ("19910630 201801 2 006", "drg. Maya Kusuma", "Dokter Gigi", "ruang-klaster-gigi", "0812-3456-7806", "PNS"),
    ("19931117 202001 1 007", "Ttg. Rizki Ramadhan, S.Kep", "Perawat Gigi", "ruang-klaster-gigi", "0812-3456-7807", "PPPK"),
    ("19890105 201401 2 008", "Ns. Fitri Handayani, S.Kep", "Perawat Skrining", "meja-skrining", "0812-3456-7808", "PNS"),
    ("19860723 201101 1 009", "Ns. Hendra Gunawan, S.Kep", "Perawat", "ruang-tindakan", "0812-3456-7809", "PNS"),
    ("19940819 201901 2 010", "Andi Pratama, S.Tr.Kes", "Analis Kesehatan", "ruang-laboratorium", "0812-3456-7810", "PPPK"),
    ("19910227 201601 2 011", "Apt. Nina Marlina, S.Farm", "Apoteker", "ruang-farmasi", "0812-3456-7811", "PNS"),
    ("19970411 202201 1 012", "Taufik Hidayat, A.Md.Farm", "Tenaga Teknis Kefarmasian", "ruang-farmasi", "0812-3456-7812", "PPPK"),
    ("19950902 202101 2 013", "Rina Puspita, A.Md", "Staf Administrasi", "loket", "0812-3456-7813", "Non-ASN"),
    ("19961225 202201 2 014", "Dedi Kurniawan", "Petugas Pendaftaran", "pendaftaran-antrian", "0812-3456-7814", "Non-ASN"),
    ("19830814 200901 1 015", "Irma Suryani, S.KM", "Kepala TU", "ruang-tu", "0812-3456-7815", "PNS"),
    ("19931007 201801 2 016", "Sari Wulandari, A.Md.Kes", "Epidemiolog", "klaster-4", "0812-3456-7816", "PPPK"),
]


@app.on_event("startup")
async def seed_data():
    if await db.pegawai.count_documents({}) == 0:
        now = datetime.now(timezone.utc).date()
        docs = [{
            "id": str(uuid.uuid4()), "nip": nip, "nama": nama, "jabatan": jab,
            "ruangan_id": rid, "telepon": telp, "status_kepegawaian": st,
            "created_at": datetime.now(timezone.utc).isoformat(),
        } for nip, nama, jab, rid, telp, st in SAMPLE_PEGAWAI]
        await db.pegawai.insert_many(docs)
        today = now.isoformat()
        besok = (now + timedelta(days=1)).isoformat()
        lusa = (now + timedelta(days=2)).isoformat()
        minggu_depan = (now + timedelta(days=5)).isoformat()
        j1 = {
            "id": str(uuid.uuid4()), "nama_kegiatan": "Posyandu Balita Kelurahan Sukamaju",
            "tanggal_mulai": today, "tanggal_selesai": today,
            "lokasi": "Balai RW 05 Sukamaju", "keterangan": "Penimbangan, imunisasi dasar, dan penyuluhan gizi",
            "pegawai_ids": [docs[3]["id"], docs[4]["id"]], "status": "disetujui",
            "disetujui_oleh": "Kepala Puskesmas", "catatan_approval": "", "waktu_approval": datetime.now(timezone.utc).isoformat(),
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        j2 = {
            "id": str(uuid.uuid4()), "nama_kegiatan": "Puskesmas Keliling (Pusling) Desa Mekarwangi",
            "tanggal_mulai": besok, "tanggal_selesai": lusa,
            "lokasi": "Desa Mekarwangi", "keterangan": "Pelayanan pemeriksaan umum dan pengobatan gratis",
            "pegawai_ids": [docs[0]["id"], docs[8]["id"], docs[10]["id"]], "status": "menunggu",
            "disetujui_oleh": "", "catatan_approval": "", "waktu_approval": "",
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        j3 = {
            "id": str(uuid.uuid4()), "nama_kegiatan": "Bulan Imunisasi Anak Sekolah (BIAS) SDN 01",
            "tanggal_mulai": minggu_depan, "tanggal_selesai": minggu_depan,
            "lokasi": "SDN 01 Cempaka", "keterangan": "Imunisasi DT dan TT untuk siswa kelas 1-3",
            "pegawai_ids": [docs[4]["id"], docs[2]["id"]], "status": "menunggu",
            "disetujui_oleh": "", "catatan_approval": "", "waktu_approval": "",
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        await db.jadwal.insert_many([j1, j2, j3])


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
